"""Regression checks for the manually formatted report front matter."""

from __future__ import annotations

import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.table import _Cell  # pyright: ignore[reportPrivateUsage]

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "scripts"))

from update_report import REPORT_PATH, synchronize_report  # noqa: E402


def _protected_prefix(path: Path) -> tuple[bytes, ...]:
    document = Document(str(path))
    protected: list[bytes] = [paragraph._p.xml.encode() for paragraph in document.paragraphs]

    def visit_cell(cell: _Cell) -> bool:
        for paragraph in cell.paragraphs:
            if paragraph.text.strip() == "需求分析说明书（10分）":
                return True
            protected.append(paragraph._p.xml.encode())
        for table in cell.tables:
            for row in table.rows:
                for nested_cell in row.cells:
                    if visit_cell(nested_cell):
                        return True
        return False

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if visit_cell(cell):
                    return tuple(protected)
    raise AssertionError("requirements section marker is missing")


def _protected_document(path: Path) -> bytes:
    """Mask generated chapter content and compare everything else in the Word body."""
    body = deepcopy(Document(str(path)).element.body)
    chapter_titles = {"需求分析说明书（10分）", "系统设计文档（20分）", "软件测试报告（20分）"}
    for cell in body.iter():
        if cell.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc":
            continue
        for index, child in enumerate(list(cell)):
            text = "".join(node.text or "" for node in child.iter() if node.tag.endswith("}t"))
            if text.strip() in chapter_titles:
                for generated in list(cell)[index:]:
                    cell.remove(generated)
                break
    return body.xml.encode()


class ReportPreservationTests(unittest.TestCase):
    def test_sync_preserves_cover_and_front_matter(self) -> None:
        if not REPORT_PATH.is_file():
            self.skipTest("local report.docx is ignored by Git and is not present")
        with tempfile.TemporaryDirectory() as directory:
            temporary_report = Path(directory) / "report.docx"
            copy2(REPORT_PATH, temporary_report)
            before = _protected_prefix(temporary_report)
            document_before = _protected_document(temporary_report)
            synchronize_report(temporary_report)
            self.assertEqual(_protected_prefix(temporary_report), before)
            self.assertEqual(_protected_document(temporary_report), document_before)
            document = Document(str(temporary_report))
            text = "\n".join(
                node.text or "" for node in document.element.iter() if node.tag.endswith("}t")
            )
            for required in [
                "单元测试",
                "组件集成测试",
                "稳定性测试",
                "静态代码分析",
                "Instruments",
            ]:
                self.assertIn(required, text)
            self.assertNotIn("backup_available=false", text)


if __name__ == "__main__":
    unittest.main()
