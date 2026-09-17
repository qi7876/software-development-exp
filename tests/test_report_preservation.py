"""Regression checks for the manually formatted report front matter."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.table import _Cell  # pyright: ignore[reportPrivateUsage]

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from update_report import REPORT_PATH, synchronize_report  # noqa: E402


def _protected_prefix(path: Path) -> tuple[bytes, ...]:
    document = Document(str(path))
    protected: list[bytes] = [paragraph._p.xml.encode() for paragraph in document.paragraphs]

    def visit_cell(cell: _Cell) -> bool:
        for paragraph in cell.paragraphs:
            if paragraph.text.strip() == "需求分析说明书（10分）":
                return True
            protected.append(paragraph._p.xml.encode())
        for table in cell.tables:
            for row in table.rows:
                for nested_cell in row.cells:
                    if visit_cell(nested_cell):
                        return True
        return False

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if visit_cell(cell):
                    return tuple(protected)
    raise AssertionError("requirements section marker is missing")


class ReportPreservationTests(unittest.TestCase):
    def test_sync_preserves_cover_and_front_matter(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            temporary_report = Path(directory) / "report.docx"
            copy2(REPORT_PATH, temporary_report)
            before = _protected_prefix(temporary_report)
            synchronize_report(temporary_report)
            self.assertEqual(_protected_prefix(temporary_report), before)


if __name__ == "__main__":
    unittest.main()
